from base64 import b64encode
from dataclasses import asdict
from io import BytesIO
from json import dumps
from pathlib import Path
from random import Random
from subprocess import CompletedProcess
from subprocess import run as run_process
from tempfile import TemporaryDirectory
from typing import Any, BinaryIO

from botocore.exceptions import ClientError
from docx import Document as DocxDocument
from fakes.fake_uploadcare_client import FakeUploadcareFile
from fastapi.security import HTTPAuthorizationCredentials
from google.genai.types import GenerateContentResponse, GroundingChunk, Model
from langchain_core.messages import AIMessage, HumanMessage
from langchain_core.outputs import ChatResult
from openai.types import CreateEmbeddingResponse, Embedding
from openai.types.audio import Transcription
from perplexity.types.shared.api_public_search_result import APIPublicSearchResult
from PIL import Image
from replicate.client import Client as ReplicateClient
from replicate.helpers import FileOutput
from replicate.prediction import Prediction
from requests import Response
from urllib3.response import HTTPResponse
from xai_sdk.chat import Response as XAIChatResponse
from xai_sdk.proto import chat_pb2, image_pb2
from xai_sdk.sync.image import ImageResponse as XAIImageResponse

from features.chat.telegram.model.attachment.audio import Audio as TelegramAudio
from features.chat.telegram.model.attachment.document import Document as TelegramDocument
from features.chat.telegram.model.attachment.photo_size import PhotoSize as TelegramPhotoSize
from features.chat.telegram.model.attachment.video import Video as TelegramVideo
from features.chat.telegram.model.attachment.voice import Voice as TelegramVoice
from features.chat.telegram.model.chat import Chat as TelegramChat
from features.chat.telegram.model.chat_member import (
    ChatMemberAdministrator,
    ChatMemberBanned,
    ChatMemberLeft,
    ChatMemberMember,
    ChatMemberOwner,
    ChatMemberRestricted,
)
from features.chat.telegram.model.message import Message as TelegramMessage
from features.chat.telegram.model.text_quote import TextQuote as TelegramTextQuote
from features.chat.telegram.model.update import Update as TelegramUpdate
from features.chat.telegram.model.user import User as TelegramUser
from features.chat.whatsapp.model.attachment.media_attachment import MediaAttachment as WhatsAppMediaAttachment
from features.chat.whatsapp.model.attachment.text import Text as WhatsAppText
from features.chat.whatsapp.model.change import Change as WhatsAppChange
from features.chat.whatsapp.model.contact import Contact as WhatsAppContact
from features.chat.whatsapp.model.context import Context as WhatsAppContext
from features.chat.whatsapp.model.entry import Entry as WhatsAppEntry
from features.chat.whatsapp.model.message import Message as WhatsAppMessage
from features.chat.whatsapp.model.metadata import Metadata as WhatsAppMetadata
from features.chat.whatsapp.model.profile import Profile as WhatsAppProfile
from features.chat.whatsapp.model.response import (
    ContactResponse as WhatsAppContactResponse,
)
from features.chat.whatsapp.model.response import (
    MarkAsReadResponse as WhatsAppMarkAsReadResponse,
)
from features.chat.whatsapp.model.response import (
    MessageResponse as WhatsAppMessageResponse,
)
from features.chat.whatsapp.model.response import (
    SentMessageResponse as WhatsAppSentMessageResponse,
)
from features.chat.whatsapp.model.status import MessageStatus as WhatsAppMessageStatus
from features.chat.whatsapp.model.status import StatusPricing as WhatsAppStatusPricing
from features.chat.whatsapp.model.update import Update as WhatsAppUpdate
from features.chat.whatsapp.model.value import Value as WhatsAppValue
from features.web_browsing.twitter_status_fetcher import (
    TweetData,
    TweetLinkPreview,
    TweetMediaItem,
    TweetMediaVariant,
    TweetUserData,
)


def ai_message(**overrides: Any) -> AIMessage:
    defaults = {"content": "Hello world"}
    return AIMessage(**(defaults | overrides))


def human_message(**overrides: Any) -> HumanMessage:
    defaults = {"content": "Hello world"}
    return HumanMessage(**(defaults | overrides))


def chat_result(**overrides: Any) -> ChatResult:
    return ChatResult(**({"generations": [{"message": ai_message()}]} | overrides))


def google_generate_content_response(**overrides: Any) -> GenerateContentResponse:
    defaults = {
        "usage_metadata": {
            "prompt_token_count": 100,
            "candidates_token_count": 200,
            "total_token_count": 300,
        },
    }
    return GenerateContentResponse(**(defaults | overrides))


def google_grounding_response(
    query_count: int = 2,
    text: str = "Google answer",
    grounding_chunks: list[GroundingChunk] | None = None,
    **overrides: Any,
) -> GenerateContentResponse:
    defaults = {
        "usage_metadata": {
            "prompt_token_count": 10,
            "candidates_token_count": 200,
            "thoughts_token_count": 50,
            "total_token_count": 260,
        },
        "candidates": [{
            "content": {"parts": [{"text": text}], "role": "model"},
            "grounding_metadata": {
                "web_search_queries": [f"query {i}" for i in range(query_count)],
                "grounding_chunks": grounding_chunks or [],
            },
        }],
    }
    return google_generate_content_response(**(defaults | overrides))


def google_model(**overrides: Any) -> Model:
    return Model(**({"name": "test-model", "display_name": "Test Model"} | overrides))


def google_grounding_chunk(**overrides: Any) -> GroundingChunk:
    defaults = {"web": {"title": "example.com", "uri": "https://example.com/page"}}
    return GroundingChunk(**(defaults | overrides))


def perplexity_search_result(**overrides: Any) -> APIPublicSearchResult:
    defaults = {"title": "Example page", "url": "https://example.com/page"}
    return APIPublicSearchResult(**(defaults | overrides))


def x_ai_chat_response(content: str = "xAI answer", **overrides: Any) -> XAIChatResponse:
    defaults = {
        "id": "response-123",
        "model": "grok-4.3",
        "outputs": [{"message": {"role": "ROLE_ASSISTANT", "content": content}}],
        "usage": {
            "cost_in_usd_ticks": 25_000_000,
            "prompt_tokens": 10,
            "completion_tokens": 20,
            "total_tokens": 30,
        },
    }
    return XAIChatResponse(chat_pb2.GetChatCompletionResponse(**(defaults | overrides)), index = None)


def x_ai_image_response(content: bytes | None = None, **overrides: Any) -> XAIImageResponse:
    images = [{"url": "https://example.com/image.png"}] if content is None else [
        {"base64": "data:image/png;base64," + b64encode(content).decode(), "respect_moderation": True},
    ]
    defaults = {"model": "grok-imagine-image", "images": images}
    return XAIImageResponse(image_pb2.ImageResponse(**(defaults | overrides)), index = 0)


def openai_transcription(**overrides: Any) -> Transcription:
    return Transcription(**({"text": "Audio transcription"} | overrides))


def openai_embedding(**overrides: Any) -> Embedding:
    return Embedding(**({"embedding": [1.0, 0.0], "index": 0, "object": "embedding"} | overrides))


def openai_embedding_response(**overrides: Any) -> CreateEmbeddingResponse:
    defaults = {
        "data": [openai_embedding()],
        "model": "text-embedding-3-small",
        "object": "list",
        "usage": {"prompt_tokens": 1, "total_tokens": 1},
    }
    return CreateEmbeddingResponse(**(defaults | overrides))


def process_result(**overrides: Any) -> CompletedProcess[str]:
    defaults = {"args": [], "returncode": 0, "stdout": "", "stderr": ""}
    return CompletedProcess(**(defaults | overrides))


def image_bitmap(
    size: tuple[int, int] = (16, 16),
    color: tuple[int, int, int] | tuple[int, int, int, int] = (100, 150, 200),
    noisy: bool = False,
) -> Image.Image:
    mode = "RGBA" if len(color) == 4 else "RGB"
    if noisy:
        return Image.frombytes(mode, size, Random(42).randbytes(size[0] * size[1] * len(color)))
    return Image.new(mode, size, color = color)


def image_bytes(
    image_format: str = "PNG",
    size: tuple[int, int] = (16, 16),
    color: tuple[int, int, int] | tuple[int, int, int, int] = (100, 150, 200),
    **save_options: Any,
) -> bytes:
    buffer = BytesIO()
    with image_bitmap(size = size, color = color) as image:
        image.save(buffer, format = image_format, **save_options)
    return buffer.getvalue()


def image_file(
    path: Path,
    size: tuple[int, int] = (16, 16),
    color: tuple[int, int, int] | tuple[int, int, int, int] = (100, 150, 200),
) -> Path:
    path.write_bytes(image_bytes(size = size, color = color))
    return path


def svg_bytes() -> bytes:
    return b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"><circle cx="5" cy="5" r="4" fill="white"/></svg>'


def telegram_chat_member(**overrides: Any) -> ChatMemberMember:
    defaults = {"status": "member", "user": telegram_user()}
    return ChatMemberMember(**(defaults | overrides))


def telegram_chat_owner(**overrides: Any) -> ChatMemberOwner:
    defaults = {"status": "creator", "user": telegram_user(), "is_anonymous": False}
    return ChatMemberOwner(**(defaults | overrides))


def telegram_chat_administrator(**overrides: Any) -> ChatMemberAdministrator:
    defaults = {
        "status": "administrator", "user": telegram_user(),
        "can_be_edited": False, "is_anonymous": False, "can_manage_chat": True,
        "can_delete_messages": False, "can_manage_video_chats": False,
        "can_restrict_members": False, "can_promote_members": False,
        "can_change_info": False, "can_invite_users": False,
        "can_post_stories": False, "can_edit_stories": False, "can_delete_stories": False,
    }
    return ChatMemberAdministrator(**(defaults | overrides))


def telegram_chat_member_restricted(**overrides: Any) -> ChatMemberRestricted:
    defaults = {
        "status": "restricted", "user": telegram_user(), "is_member": True,
        "can_send_messages": False, "can_send_audios": False, "can_send_documents": False,
        "can_send_photos": False, "can_send_videos": False, "can_send_video_notes": False,
        "can_send_voice_notes": False, "can_send_polls": False, "can_send_other_messages": False,
        "can_add_web_page_previews": False, "can_change_info": False, "can_invite_users": False,
        "can_pin_messages": False, "can_manage_topics": False, "until_date": 0,
    }
    return ChatMemberRestricted(**(defaults | overrides))


def telegram_chat_member_banned(**overrides: Any) -> ChatMemberBanned:
    defaults = {"status": "kicked", "user": telegram_user(), "until_date": 0}
    return ChatMemberBanned(**(defaults | overrides))


def telegram_chat_member_left(**overrides: Any) -> ChatMemberLeft:
    defaults = {"status": "left", "user": telegram_user()}
    return ChatMemberLeft(**(defaults | overrides))


def telegram_message_response(**message_overrides: Any) -> dict[str, Any]:
    return {"ok": True, "result": telegram_message(**message_overrides).model_dump(by_alias = True)}


def whatsapp_mark_as_read_response(**overrides: Any) -> WhatsAppMarkAsReadResponse:
    return WhatsAppMarkAsReadResponse(**({"success": True} | overrides))


def replicate_prediction(**overrides: Any) -> Prediction:
    defaults = {
        "id": "prediction-123",
        "model": "test/model",
        "version": "test-version",
        "status": "succeeded",
        "output": None,
        "error": None,
        "logs": None,
    }
    return Prediction(**(defaults | overrides))


def replicate_file_output(url: str = "https://example.com/image.png") -> FileOutput:
    return FileOutput(url = url, client = ReplicateClient(api_token = "test-token"))


def docx_document_bytes(text: str = "A test paragraph.") -> bytes:
    document = DocxDocument()
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def http_response(
    content: bytes = b"content",
    status_code: int = 200,
    url: str = "https://example.com",
    encoding: str | None = "utf-8",
    headers: dict[str, str] | None = None,
) -> Response:
    response = Response()
    response.status_code = status_code
    response.url = url
    response.encoding = encoding
    response.headers.update(headers or {})
    response.raw = HTTPResponse(body = BytesIO(content), preload_content = False)
    return response


def http_json_response(payload: Any, status_code: int = 200, url: str = "https://example.com") -> Response:
    response = http_response(content = dumps(payload).encode("utf-8"), status_code = status_code, url = url)
    response.headers["Content-Type"] = "application/json"
    return response


def uploadcare_file(**overrides: Any) -> FakeUploadcareFile:
    defaults = {"cdn_url": "https://cdn-id.ucarecd.net/uuid/", "filename": "attachment-id.txt"}
    return FakeUploadcareFile(**(defaults | overrides))


def s3_client_error(code: str = "AccessDenied", status: int = 403, operation: str = "HeadBucket") -> ClientError:
    return ClientError({"Error": {"Code": code}, "ResponseMetadata": {"HTTPStatusCode": status}}, operation)


def s3_object_response(body: BinaryIO) -> dict[str, object]:
    return {"Body": body}


def whatsapp_profile(**overrides: Any) -> WhatsAppProfile:
    defaults = {
        "name": "New User",
    }
    return WhatsAppProfile(**(defaults | overrides))


def whatsapp_contact(**overrides: Any) -> WhatsAppContact:
    defaults = {
        "profile": whatsapp_profile(),
        "wa_id": "1",
    }
    return WhatsAppContact(**(defaults | overrides))


def whatsapp_metadata(**overrides: Any) -> WhatsAppMetadata:
    defaults = {
        "display_phone_number": "123",
        "phone_number_id": "phone-id",
    }
    return WhatsAppMetadata(**(defaults | overrides))


def whatsapp_status_pricing(**overrides: Any) -> WhatsAppStatusPricing:
    defaults = {
        "billable": True,
        "type": "regular",
        "category": "service",
    }
    return WhatsAppStatusPricing(**(defaults | overrides))


def whatsapp_message_status(**overrides: Any) -> WhatsAppMessageStatus:
    defaults = {
        "id": "whatsapp-message-123",
        "status": "sent",
        "timestamp": "1768478400",
        "recipient_id": "15551234567",
        "pricing": whatsapp_status_pricing(),
    }
    return WhatsAppMessageStatus(**(defaults | overrides))


def whatsapp_text(**overrides: Any) -> WhatsAppText:
    defaults = {
        "body": "Hello from Mark Johnson.",
    }
    return WhatsAppText(**(defaults | overrides))


def whatsapp_context(**overrides: Any) -> WhatsAppContext:
    defaults = {
        "from": None,
        "id": "message-122",
    }
    return WhatsAppContext(**(defaults | overrides))


def whatsapp_media_attachment(**overrides: Any) -> WhatsAppMediaAttachment:
    defaults = {
        "id": "whatsapp-media-123",
        "caption": None,
        "mime_type": "image/jpeg",
        "sha256": None,
        "filename": None,
        "voice": None,
    }
    return WhatsAppMediaAttachment(**(defaults | overrides))


def whatsapp_message(**overrides: Any) -> WhatsAppMessage:
    defaults: dict[str, Any] = {
        "from": "15551234567",
        "id": "whatsapp-message-123",
        "timestamp": "1768478400",
        "type": "text",
    }
    if "text" not in overrides:
        defaults["text"] = whatsapp_text()
    return WhatsAppMessage(**(defaults | overrides))


def whatsapp_value(**overrides: Any) -> WhatsAppValue:
    defaults = {
        "messaging_product": "whatsapp",
        "metadata": whatsapp_metadata(),
        "contacts": [whatsapp_contact()],
        "messages": [],
    }
    return WhatsAppValue(**(defaults | overrides))


def whatsapp_change(**overrides: Any) -> WhatsAppChange:
    defaults: dict[str, Any] = {
        "field": "messages",
    }
    if "value" not in overrides:
        defaults["value"] = whatsapp_value()
    return WhatsAppChange(**(defaults | overrides))


def whatsapp_entry(**overrides: Any) -> WhatsAppEntry:
    defaults: dict[str, Any] = {
        "id": "whatsapp-business-account-123",
    }
    if "changes" not in overrides:
        defaults["changes"] = [whatsapp_change()]
    return WhatsAppEntry(**(defaults | overrides))


def whatsapp_update(**overrides: Any) -> WhatsAppUpdate:
    defaults: dict[str, Any] = {
        "object": "whatsapp_business_account",
    }
    if "entry" not in overrides:
        defaults["entry"] = [whatsapp_entry()]
    return WhatsAppUpdate(**(defaults | overrides))


def whatsapp_contact_response(**overrides: Any) -> WhatsAppContactResponse:
    defaults = {
        "input": "+15551234567",
        "wa_id": "15551234567",
    }
    return WhatsAppContactResponse(**(defaults | overrides))


def whatsapp_sent_message_response(**overrides: Any) -> WhatsAppSentMessageResponse:
    defaults = {
        "id": "whatsapp-message-123",
        "message_status": "accepted",
    }
    return WhatsAppSentMessageResponse(**(defaults | overrides))


def whatsapp_message_response(**overrides: Any) -> WhatsAppMessageResponse:
    defaults: dict[str, Any] = {
        "messaging_product": "whatsapp",
    }
    if "contacts" not in overrides:
        defaults["contacts"] = [whatsapp_contact_response()]
    if "messages" not in overrides:
        defaults["messages"] = [whatsapp_sent_message_response()]
    return WhatsAppMessageResponse(**(defaults | overrides))


def telegram_chat(**overrides: Any) -> TelegramChat:
    defaults = {
        "id": 123456789,
        "type": "private",
        "title": None,
        "username": None,
        "first_name": None,
        "last_name": None,
    }
    return TelegramChat(**(defaults | overrides))


def telegram_user(**overrides: Any) -> TelegramUser:
    defaults = {
        "id": 123456789,
        "is_bot": False,
        "first_name": "Mark",
        "last_name": None,
        "username": None,
        "language_code": None,
    }
    return TelegramUser(**(defaults | overrides))


def telegram_audio(**overrides: Any) -> TelegramAudio:
    defaults = {
        "file_id": "telegram-audio-123",
        "file_unique_id": "telegram-unique-audio-123",
        "file_name": None,
        "mime_type": None,
        "file_size": None,
    }
    return TelegramAudio(**(defaults | overrides))


def telegram_document(**overrides: Any) -> TelegramDocument:
    defaults = {
        "file_id": "telegram-file-123",
        "file_unique_id": "telegram-unique-file-123",
        "file_name": None,
        "mime_type": None,
        "file_size": None,
    }
    return TelegramDocument(**(defaults | overrides))


def telegram_voice(**overrides: Any) -> TelegramVoice:
    defaults = {
        "file_id": "telegram-voice-123",
        "file_unique_id": "telegram-unique-voice-123",
        "mime_type": None,
        "file_size": None,
    }
    return TelegramVoice(**(defaults | overrides))


def telegram_video(**overrides: Any) -> TelegramVideo:
    defaults = {
        "file_id": "telegram-video-123",
        "file_unique_id": "telegram-unique-video-123",
        "width": 1920,
        "height": 1080,
        "duration": 12,
        "file_name": None,
        "mime_type": None,
        "file_size": None,
    }
    return TelegramVideo(**(defaults | overrides))


def telegram_photo_size(**overrides: Any) -> TelegramPhotoSize:
    defaults = {
        "file_id": "telegram-photo-123",
        "file_unique_id": "telegram-unique-photo-123",
        "width": 1280,
        "height": 720,
        "file_size": None,
    }
    return TelegramPhotoSize(**(defaults | overrides))


def telegram_text_quote(**overrides: Any) -> TelegramTextQuote:
    defaults = {
        "text": "Selected quote",
        "position": 0,
        "entities": None,
    }
    return TelegramTextQuote(**(defaults | overrides))


def telegram_message(**overrides: Any) -> TelegramMessage:
    defaults: dict[str, Any] = {
        "message_id": 123,
        "text": None,
        "date": 1_768_478_400,
    }
    if "chat" not in overrides:
        defaults["chat"] = telegram_chat()
    return TelegramMessage(**(defaults | overrides))


def telegram_update(**overrides: Any) -> TelegramUpdate:
    defaults = {
        "update_id": 123,
    }
    return TelegramUpdate(**(defaults | overrides))


def http_authorization_credentials(**overrides: Any) -> HTTPAuthorizationCredentials:
    defaults = {
        "scheme": "Bearer",
        "credentials": "valid-token",
    }
    return HTTPAuthorizationCredentials(**(defaults | overrides))


def x_tweet_response(tweet: TweetData | None = None, **overrides: Any) -> dict[str, Any]:
    defaults = {
        "data": {"text": "Test tweet content", "lang": "en", "author_id": "123"},
        "includes": {
            "users": [{
                "id": "123",
                "username": "testuser",
                "name": "Test User",
                "description": "Test bio",
            }],
        },
    }
    if tweet is not None:
        defaults["data"] = {
            "text": tweet.text, "lang": tweet.language, "created_at": tweet.created_at, "author_id": "123",
            "entities": {"urls": [
                {
                    "url": link.expanded_url, "expanded_url": link.expanded_url,
                    "title": link.title, "description": link.description,
                    "images": [{"url": link.og_image_url}] if link.og_image_url else [],
                }
                for link in tweet.link_previews
            ]},
            "referenced_tweets": [
                {"type": kind, "id": identifier}
                for kind, identifier in (("quoted", tweet.quoted_tweet_id), ("replied_to", tweet.replied_to_tweet_id))
                if identifier is not None
            ],
        }
        defaults["includes"] = {
            "users": [{
                "id": "123", "name": tweet.user.name, "username": tweet.user.handle,
                "description": tweet.user.bio, "profile_image_url": tweet.user.profile_image_url,
            }],
            "media": [
                {
                    "url": media.url, "preview_image_url": media.preview_url, "type": media.media_type,
                    "variants": [asdict(variant) for variant in media.variants],
                    "duration_ms": media.duration_ms, "width": media.width, "height": media.height,
                    "alt_text": media.alt_text,
                }
                for media in tweet.media
            ],
        }
    return defaults | overrides


def tweet_media_variant(**overrides: Any) -> TweetMediaVariant:
    defaults = {
        "url": "https://video.twimg.com/video.mp4",
        "content_type": "video/mp4",
        "bit_rate": 2_176_000,
    }
    return TweetMediaVariant(**(defaults | overrides))


def tweet_media_item(**overrides: Any) -> TweetMediaItem:
    defaults = {
        "url": "https://pbs.twimg.com/media/photo.jpg",
        "preview_url": "https://pbs.twimg.com/media/photo-preview.jpg",
        "media_type": "photo",
        "variants": [],
        "duration_ms": None,
        "width": 1280,
        "height": 720,
        "alt_text": "A city skyline at sunset.",
    }
    return TweetMediaItem(**(defaults | overrides))


def tweet_user_data(**overrides: Any) -> TweetUserData:
    defaults = {
        "name": "Mark Johnson",
        "handle": "mark_johnson",
        "bio": "Building useful software.",
        "profile_image_url": "https://pbs.twimg.com/profile_images/mark.jpg",
    }
    return TweetUserData(**(defaults | overrides))


def tweet_link_preview(**overrides: Any) -> TweetLinkPreview:
    defaults = {
        "title": "Example story",
        "description": "A concise preview of the linked story.",
        "og_image_url": "https://example.com/articles/story.jpg",
        "expanded_url": "https://example.com/articles/story",
        "domain": "example.com",
    }
    return TweetLinkPreview(**(defaults | overrides))


def tweet_data(**overrides: Any) -> TweetData:
    defaults: dict[str, Any] = {
        "text": "A useful update from Mark Johnson.",
        "language": "en",
        "created_at": "2026-01-15T12:00:00Z",
        "quoted_tweet_id": None,
        "is_reply": False,
        "replied_to_tweet_id": None,
    }
    if "user" not in overrides:
        defaults["user"] = tweet_user_data()
    if "media" not in overrides:
        defaults["media"] = [tweet_media_item()]
    if "link_previews" not in overrides:
        defaults["link_previews"] = [tweet_link_preview()]
    return TweetData(**(defaults | overrides))


def video_bytes(fast_start: bool = True) -> bytes:
    with TemporaryDirectory() as directory:
        path = Path(directory) / "video.mp4"
        run_process(
            [
                "ffmpeg", "-v", "error", "-y", "-f", "lavfi",
                "-i", "color=c=blue:s=160x90:r=10", "-t", "0.2",
                "-c:v", "libx264", "-pix_fmt", "yuv420p",
                *(["-movflags", "+faststart"] if fast_start else []), str(path),
            ],
            capture_output = True,
            check = True,
        )
        return path.read_bytes()


def iso_media_box(box_type: bytes, payload: bytes = b"") -> bytes:
    return (8 + len(payload)).to_bytes(4, byteorder = "big") + box_type + payload


def fiat_exchange_response(currency: str = "EUR", rate: float = 0.85) -> dict[str, Any]:
    return {"rates": {currency: {"rate_for_amount": str(rate)}}}


def crypto_exchange_response(symbol: str = "BTC", price: float = 40_000) -> dict[str, Any]:
    return {"data": {symbol: {"quote": {"USD": {"price": price}}}}}


def stock_quote_response(**overrides: Any) -> dict[str, Any]:
    defaults = {
        "symbol": "AAPL",
        "currency": "USD",
        "close": "210.5",
        "timestamp": 1_753_352_400,
        "is_market_open": True,
        "name": "Apple Inc.",
        "exchange": "NASDAQ",
        "mic_code": "XNAS",
        "previous_close": "208.5",
        "change": "2.0",
        "percent_change": "0.96",
    }
    return defaults | overrides


def stock_quote_error_response(code: int = 500, message: str = "Price unavailable") -> dict[str, Any]:
    return {"status": "error", "code": code, "message": message}


def github_issue_response(**overrides: Any) -> dict[str, Any]:
    return {"html_url": "https://github.com/appifyhub/agent-backend/issues/123"} | overrides
